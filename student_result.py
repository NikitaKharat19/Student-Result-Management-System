from openpyxl import Workbook, load_workbook
import os


file_name = "student_results.xlsx"


def create_file():
    if not os.path.exists(file_name):
        wb = Workbook()
        ws = wb.active
        ws.title = "Student Results"

        ws.append([
            "Roll No", "Name", "Class",
            "Subject 1", "Subject 2", "Subject 3",
            "Subject 4", "Subject 5",
            "Total", "Percentage", "Grade", "Status"
        ])

        wb.save(file_name)


def calculate_result(marks):
    total = sum(marks)
    percentage = total / 5

    if any(mark < 40 for mark in marks):
        status = "FAIL"
        grade = "F"
    elif percentage >= 90:
        grade = "A+"
        status = "PASS"
    elif percentage >= 75:
        grade = "A"
        status = "PASS"
    elif percentage >= 60:
        grade = "B"
        status = "PASS"
    elif percentage >= 40:
        grade = "C"
        status = "PASS"
    else:
        grade = "F"
        status = "FAIL"

    return total, percentage, grade, status


def add_student():
    roll_no = int(input("enter roll no: "))
    name = input("enter student name: ")
    student_class = input("enter class: ")

    marks = []

    for i in range(1, 6):
        mark = int(input("enter marks of subject " + str(i) + ": "))
        marks.append(mark)

    total, percentage, grade, status = calculate_result(marks)

    wb = load_workbook(file_name)
    ws = wb.active

    ws.append([
        roll_no,
        name,
        student_class,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        percentage,
        grade,
        status
    ])

    wb.save(file_name)

    print()
    print("student result added successfully")
    print("total =", total)
    print("percentage =", percentage)
    print("grade =", grade)
    print("status =", status)


def get_result():
    roll_no = int(input("enter roll no: "))

    wb = load_workbook(file_name)
    ws = wb.active

    found = False

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == roll_no:
            print()
            print("--------------------------------")
            print("student result")
            print("--------------------------------")
            print("roll no     :", row[0])
            print("name        :", row[1])
            print("class       :", row[2])
            print("total       :", row[8])
            print("percentage  :", format(row[9], ".2f") + "%")
            print("grade       :", row[10])
            print("status      :", row[11])
            print("--------------------------------")

            found = True
            break

    if not found:
        print("student not found")


def show_all_data():
    wb = load_workbook(file_name)
    ws = wb.active

    print()
    print("student result data")
    print("-" * 85)

    print(
        "roll no\tname\tclass\ttotal\tpercentage\tgrade\tstatus"
    )

    print("-" * 85)

    for row in ws.iter_rows(min_row=2, values_only=True):
        print(
            row[0], "\t",
            row[1], "\t",
            row[2], "\t",
            row[8], "\t",
            format(row[9], ".2f"), "\t\t",
            row[10], "\t",
            row[11]
        )

    print("-" * 85)


def menu():
    create_file()

    while True:
        print()
        print("========================================")
        print("      student result management")
        print("========================================")
        print("1. add student result")
        print("2. get student result")
        print("3. show all student data")
        print("4. exit")
        print("========================================")

        choice = int(input("enter your choice: "))

        if choice == 1:
            add_student()

        elif choice == 2:
            get_result()

        elif choice == 3:
            show_all_data()

        elif choice == 4:
            print("thank you")
            break

        else:
            print("invalid choice")


menu()